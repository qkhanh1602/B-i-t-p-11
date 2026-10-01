import sys, os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text_list, title=None, fill_hex="F0F4F8", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(title + "\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    for i, line in enumerate(text_list):
        if i > 0 or title:
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_before = Pt(1)
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.line_spacing = 1.15
            run = p2.add_run(line)
        else:
            run = p.add_run(line)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, title, code_snippet, max_lines=45):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r_icon = p_title.add_run("💻 Code: ")
    r_icon.bold = True
    r_icon.font.name = "Calibri"
    r_icon.font.size = Pt(10.5)
    r_icon.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)
    
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.name = "Consolas"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(0x8E, 0x44, 0xAD)

    lines = code_snippet.strip().splitlines()
    if len(lines) > max_lines:
        snippet_text = "\n".join(lines[:max_lines]) + f"\n... [đã rút gọn {len(lines) - max_lines} dòng tiếp theo]"
    else:
        snippet_text = "\n".join(lines)
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="4A90E2"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(snippet_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(6)
    r_line = p_line.add_run("―" * 55)
    r_line.font.color.rgb = RGBColor(0xBD, 0xC3, 0xC7)
    r_line.font.size = Pt(8)

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)

def add_image_figure(doc, img_path, caption):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found.")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(6.2))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(1)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap_lbl = p_cap.add_run("Hình: ")
    r_cap_lbl.bold = True
    r_cap_lbl.font.italic = True
    r_cap_lbl.font.size = Pt(9.5)
    r_cap_lbl.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    r_cap = p_cap.add_run(caption)
    r_cap.font.italic = True
    r_cap.font.size = Pt(9.5)
    r_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

print("Helper module loaded successfully.")
